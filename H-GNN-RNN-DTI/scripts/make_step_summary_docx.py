"""Generate a plain-English Word summary of the step 1 & 2 artifacts.

Output: docs/Steps1-2_Summary_PlainEnglish.docx
Run:    python scripts/make_step_summary_docx.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def h(doc, text, level):
    p = doc.add_heading(text, level=level)
    return p


def para(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    try:
        t.style = "Light Grid Accent 1"
    except KeyError:
        t.style = "Table Grid"
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(htext)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return t


def build(out_path: Path) -> None:
    doc = Document()

    # Title
    title = doc.add_heading("Drug–Target AI Model — Plain-English Summary of the Data Setup", level=0)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = sub.add_run("Objective 1 (the new “fusion” mechanism)  •  Steps 1 & 2: getting and preparing the data  •  17 October 2026")
    r.italic = True
    r.font.color.rgb = ACCENT

    # What this is
    h(doc, "What this document is", 1)
    para(doc,
         "We have started building the AI model for Objective 1. Before going further "
         "(building the model itself), here is a plain-English summary of what has been done "
         "so far — collecting the data and preparing it — so you can review and confirm we are "
         "on the right track.")

    # Big picture
    h(doc, "The big picture (what we are building)", 1)
    para(doc,
         "The model predicts whether a drug will interact with a protein target in the body. "
         "A drug is a small molecule; a protein is a long chain of building blocks called amino acids. "
         "The model learns patterns from thousands of known, lab-measured examples.")
    para(doc,
         "Objective 1 is specifically about a new way of letting the drug information and the protein "
         "information “talk to each other” inside the model (we call this the fusion mechanism). Before "
         "we can test whether that new idea works better than existing approaches, we need good data, and "
         "the data must be put into a form the model can read. That is exactly what Steps 1 and 2 did.")

    # Step 1
    h(doc, "Step 1 — Getting the data (done)", 1)
    para(doc,
         "We downloaded standard, widely used public datasets of measured drug–protein pairs. "
         "Each row is one drug paired with one protein, plus a number measured in a laboratory that says "
         "how strongly they stick together.")
    table(doc,
          ["Dataset", "What it is", "Pairs", "Drugs", "Proteins", "The measured number"],
          [
              ["DAVIS", "Cancer-related (kinase) binding strengths — our starting dataset", "25,772", "68", "379", "Kd (lower = stronger binding)"],
              ["KIBA", "A larger kinase dataset using a combined score", "117,657", "2,068", "229", "KIBA score (higher = stronger)"],
              ["BindingDB (Kd)", "A broad public binding database", "52,274", "10,636", "1,090", "Kd (lower = stronger binding)"],
          ])
    para(doc, "Notes:", bold=True)
    bullet(doc, "We start with DAVIS because it is the smallest and cleanest, so we can get the whole pipeline working quickly before scaling up.")
    bullet(doc, "Two older datasets mentioned in the proposal (Yamanishi and Luo) are not available from this same source. We recorded exactly where to obtain them for later, rather than skipping them silently.")
    bullet(doc, "BindingDB contains a few impossible “0” binding values (a known data-quality quirk). We will clean those before using BindingDB. DAVIS has no such problem.")

    # Step 2
    h(doc, "Step 2 — Preparing the data (done)", 1)
    para(doc, "Raw data is not something the model can read directly. We transformed it in four ways:")

    para(doc, "1. Turning each drug molecule into a “graph.”", bold=True)
    para(doc,
         "A molecule is naturally a set of atoms joined by bonds — like dots connected by lines. "
         "We converted each drug’s chemical shorthand into this dots-and-lines form, with descriptive "
         "tags on every atom and bond. This is the format the drug-reading part of the model needs.")

    para(doc, "2. Turning each protein into a sequence of numbers.", bold=True)
    para(doc,
         "A protein is a long string of amino-acid letters. We converted each protein into a numbered "
         "sequence (each letter becomes a number), capped at 1,000 letters. This is the format the "
         "protein-reading part of the model needs.")

    para(doc, "3. Deciding “interacts” or “does not interact.”", bold=True)
    para(doc,
         "We applied the agreed rule: if the measured binding strength is 30 nM or stronger "
         "(written Kd ≤ 30; smaller numbers mean stronger binding), the pair is labelled “interacts (yes),” "
         "otherwise “no.” In DAVIS, only about 4% of pairs are a “yes.” This imbalance is expected and is "
         "the very reason we use a special training technique (called focal loss) that stops the model from "
         "lazily always guessing “no.”")

    para(doc, "4. Splitting the data fairly, with no cheating.", bold=True)
    para(doc,
         "We divided the data into a training set (80%), a validation set (10%), and a test set (10%). "
         "We automatically checked that no exact drug–protein pair appears in more than one set, so the "
         "model cannot simply memorise a test answer. Each set keeps the same ~4% “yes” rate.")

    # Files produced
    h(doc, "What was produced (the files)", 1)
    table(doc,
          ["File", "In plain terms"],
          [
              ["data/raw/*.csv", "The raw downloaded data for each dataset."],
              ["data/processed/davis/drug_graphs.pkl", "Every DAVIS drug stored as a dots-and-lines graph."],
              ["data/processed/davis/protein_tokens.pkl", "Every DAVIS protein stored as a number sequence."],
              ["data/processed/davis/pairs.csv", "The list of pairs with their yes/no label and which set (train/val/test) they belong to."],
              ["data/processed/davis/meta.json", "A one-page summary sheet: counts, the 30 nM setting, and sizes."],
              ["data/splits/davis/S1_seed42.json", "The exact record of which rows are training, validation, and test."],
              ["docs/objective1_analysis.md", "The full technical analysis and build plan (for the record)."],
          ])

    # Trust
    h(doc, "Can we trust it? (the checks we ran)", 1)
    bullet(doc, "The dataset sizes matched the published reference exactly (DAVIS: 25,772 pairs, 68 drugs, 379 proteins).")
    bullet(doc, "No data leakage: confirmed no drug–protein pair appears in more than one of train/validation/test.")
    bullet(doc, "The imbalance (about 4% “yes”) was measured and is handled by design.")
    bullet(doc, "7 of 7 automated checks passed (molecule conversion, labelling, and fair splitting).")
    bullet(doc, "All saved files were reloaded successfully to confirm they are intact.")
    bullet(doc, "Reproducible: a fixed random seed and pinned software versions mean anyone can reproduce the exact same result.")

    # Next
    h(doc, "What happens next", 1)
    para(doc,
         "The next step is to build the model itself: the drug-graph reader and the protein reader, then "
         "the new fusion mechanism that is the heart of Objective 1, and finally the training and the "
         "head-to-head comparison of fusion methods. That comparison is the main result for Objective 1.")

    # Decisions
    h(doc, "A few things to note or confirm", 1)
    bullet(doc, "The “30 nM” rule for calling a pair interacting is applied and is easily changeable — please confirm it matches your intention.")
    bullet(doc, "We begin with DAVIS; KIBA and BindingDB are downloaded and ready for later.")
    bullet(doc, "No high-powered computer (GPU) is secured yet. This is fine for DAVIS at the current light setting, but one will be needed before running every dataset and comparison.")
    bullet(doc, "Reminder: completing the entire programme in a single week is not realistic from scratch. The realistic one-week target is the first genuine result on DAVIS.")

    doc.save(str(out_path))
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    out = Path(__file__).resolve().parents[1] / "docs" / "Steps1-2_Summary_PlainEnglish.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    build(out)
